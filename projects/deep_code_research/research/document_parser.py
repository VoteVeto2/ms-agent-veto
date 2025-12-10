"""
DocumentParser - Parse multiple document formats from references.zip
"""
import zipfile
from pathlib import Path
from typing import List, Any, Optional
import os
import json
import time


DEBUG_LOG_PATH = Path(__file__).resolve().parents[3] / ".cursor" / "debug.log"


class DocumentParser:
    """Parse multiple document formats from references.zip"""

    SUPPORTED_TEXT_EXTENSIONS = {'.md', '.txt', '.py', '.yaml', '.yml', '.json', '.rst', '.toml'}
    SUPPORTED_BINARY_EXTENSIONS = {'.pdf', '.docx', '.pptx'}

    def __init__(self, extract_dir: str = "./extracted_references"):
        self.extract_dir = Path(extract_dir)

    async def parse_zip(self, zip_path: str) -> List[Any]:
        """
        Extract zip or parse directory and parse all documents.

        Supported formats:
        - .pdf (use PyPDF2)
        - .md, .txt, .py, .yaml, .yml, .json (text files)
        - .docx, .pptx (use python-docx/python-pptx)

        Args:
            zip_path: Path to ZIP file or directory containing documents

        Returns:
            List of LlamaIndex Document objects
        """
        from llama_index.core import Document

        path = Path(zip_path)

        # region agent log
        try:
            with open(DEBUG_LOG_PATH, "a", encoding="utf-8") as _f:
                _f.write(json.dumps({
                    "sessionId": "debug-session",
                    "runId": "pre-fix",
                    "hypothesisId": "H1",
                    "location": "document_parser.parse_zip",
                    "message": "parse_zip_entry",
                    "data": {
                        "zip_path": str(path),
                        "cwd": os.getcwd(),
                        "exists": path.exists(),
                        "is_dir": path.is_dir(),
                        "is_file": path.is_file()
                    },
                    "timestamp": int(time.time() * 1000)
                }) + "\n")
        except Exception:
            pass
        # endregion

        # Handle directory input
        if path.is_dir():
            # region agent log
            try:
                with open(DEBUG_LOG_PATH, "a", encoding="utf-8") as _f:
                    _f.write(json.dumps({
                        "sessionId": "debug-session",
                        "runId": "pre-fix",
                        "hypothesisId": "H1",
                        "location": "document_parser.parse_zip",
                        "message": "parse_zip_directory_branch",
                        "data": {
                            "dir_path": str(path.resolve())
                        },
                        "timestamp": int(time.time() * 1000)
                    }) + "\n")
            except Exception:
                pass
            # endregion
            return await self._parse_directory(path)

        # Handle ZIP file
        # Create extract directory
        self.extract_dir.mkdir(parents=True, exist_ok=True)

        # Extract zip
        with zipfile.ZipFile(zip_path, 'r') as zf:
            zf.extractall(self.extract_dir)

        # Parse extracted files
        return await self._parse_directory(self.extract_dir)

    async def _parse_directory(self, dir_path: Path) -> List[Any]:
        """Parse all documents in a directory."""
        documents = []
        for file_path in dir_path.rglob('*'):
            if file_path.is_file():
                try:
                    file_docs = await self._parse_file(file_path)
                    documents.extend(file_docs)
                except Exception as e:
                    print(f"Warning: Failed to parse {file_path}: {e}")

        # region agent log
        try:
            with open(DEBUG_LOG_PATH, "a", encoding="utf-8") as _f:
                _f.write(json.dumps({
                    "sessionId": "debug-session",
                    "runId": "pre-fix",
                    "hypothesisId": "H1",
                    "location": "document_parser._parse_directory",
                    "message": "directory_parsed",
                    "data": {
                        "dir_path": str(dir_path.resolve()),
                        "file_count": len(documents)
                    },
                    "timestamp": int(time.time() * 1000)
                }) + "\n")
        except Exception:
            pass
        # endregion

        return documents

    async def _parse_file(self, file_path: Path) -> List[Any]:
        """Parse single file based on extension"""
        from llama_index.core import Document

        ext = file_path.suffix.lower()

        if ext in self.SUPPORTED_TEXT_EXTENSIONS:
            return await self._parse_text_file(file_path)
        elif ext == '.pdf':
            return await self._parse_pdf(file_path)
        elif ext == '.docx':
            return await self._parse_docx(file_path)
        elif ext == '.pptx':
            return await self._parse_pptx(file_path)
        else:
            # Skip unsupported files
            return []

    async def _parse_text_file(self, file_path: Path) -> List[Any]:
        """Parse text-based files"""
        from llama_index.core import Document

        encodings = ['utf-8', 'latin-1', 'cp1252']
        content = None

        for encoding in encodings:
            try:
                content = file_path.read_text(encoding=encoding)
                break
            except UnicodeDecodeError:
                continue

        if content is None:
            # Try binary read and decode with errors='ignore'
            content = file_path.read_bytes().decode('utf-8', errors='ignore')

        return [Document(
            text=content,
            metadata={
                'source': str(file_path),
                'filename': file_path.name,
                'extension': file_path.suffix
            }
        )]

    async def _parse_pdf(self, file_path: Path) -> List[Any]:
        """Parse PDF files"""
        from llama_index.core import Document

        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                from PyPDF2 import PdfReader
            except ImportError:
                print(f"Warning: PyPDF2/pypdf not installed, skipping {file_path}")
                return []

        try:
            reader = PdfReader(str(file_path))
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

            content = "\n\n".join(text_parts)
            return [Document(
                text=content,
                metadata={
                    'source': str(file_path),
                    'filename': file_path.name,
                    'extension': '.pdf',
                    'num_pages': len(reader.pages)
                }
            )]
        except Exception as e:
            print(f"Warning: Failed to parse PDF {file_path}: {e}")
            return []

    async def _parse_docx(self, file_path: Path) -> List[Any]:
        """Parse DOCX files"""
        from llama_index.core import Document

        try:
            from docx import Document as DocxDocument
        except ImportError:
            print(f"Warning: python-docx not installed, skipping {file_path}")
            return []

        try:
            doc = DocxDocument(str(file_path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs)

            return [Document(
                text=content,
                metadata={
                    'source': str(file_path),
                    'filename': file_path.name,
                    'extension': '.docx'
                }
            )]
        except Exception as e:
            print(f"Warning: Failed to parse DOCX {file_path}: {e}")
            return []

    async def _parse_pptx(self, file_path: Path) -> List[Any]:
        """Parse PPTX files"""
        from llama_index.core import Document

        try:
            from pptx import Presentation
        except ImportError:
            print(f"Warning: python-pptx not installed, skipping {file_path}")
            return []

        try:
            prs = Presentation(str(file_path))
            text_parts = []

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)
                if slide_text:
                    text_parts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_text))

            content = "\n\n".join(text_parts)
            return [Document(
                text=content,
                metadata={
                    'source': str(file_path),
                    'filename': file_path.name,
                    'extension': '.pptx',
                    'num_slides': len(prs.slides)
                }
            )]
        except Exception as e:
            print(f"Warning: Failed to parse PPTX {file_path}: {e}")
            return []
